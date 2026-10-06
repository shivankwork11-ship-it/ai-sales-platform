"""Bulk Manual Calling.

Cold-call leads that live only in ASBLVoiceBot (never in the CRM lead stores).
An admin uploads a campaign sheet; for each (phone, project) we build a normal
analytics_db.lead_profiles document (marked bulk_manual=True) with ai_calls +
call_conversations in the SAME shape the cron uses — so the EXISTING lead
detail page renders the identical AI Calls tab (metric cards, filter chips,
per-call transcript + recording). Stage & flag are manual-only (no relay),
stored on the same profile doc. These docs are excluded from the main Leads
list and shown only under the Bulk Calls tab.
"""
import re
import io
from datetime import datetime, timezone

from flask import Blueprint, jsonify, request, session

bulk_bp = Blueprint("bulk_calls", __name__)
_client = None

# ponytail: mirrors dashboard.py CALL_* buckets — keep in sync (stable taxonomy).
CALL_ANSWERED      = {"CONNECTED", "AUTO_CALLBACK", "PRE_SITE", "NOT_INTERESTED", "DROPPED", "VIRTUAL_TOUR"}
CALL_NOT_ANSWERED  = {"NO_ANSWER", "BUSY"}
CALL_NOT_CONNECTED = {"NOT_CONNECTED", "SWITCHED_OFF"}

STAGES = ("New Lead", "Lead Initiated", "Pre Site Visit", "Post Site Visit", "Not Interested")
FLAGS  = ("genuine", "spam", "test", "unclear")


def init_app(app, client):
    global _client
    _client = client
    app.register_blueprint(bulk_bp)
    try:
        client["analytics_db"]["lead_profiles"].create_index("bulk_manual", background=True)
    except Exception:
        pass


# ── helpers ───────────────────────────────────────────────────────────────────
def _l10(p):
    d = re.sub(r"\D", "", str(p or ""))
    return d[-10:] if len(d) >= 10 else ""


def _variants(p10):
    return [p10, "91" + p10, "+91" + p10]


def _norm_proj(x):
    return re.sub(r"^asbl\s+", "", str(x or "").strip().lower())


def _admin():
    return session.get("role") == "admin"


def _profiles():
    return _client["analytics_db"]["lead_profiles"]


def _lead_id(p10, projnorm):
    return "BULK-%s-%s" % (p10, projnorm or "")


def _fetch_call_data(phones10):
    """Bulk-fetch (3 queries total, not per-lead) all schedules/transcripts/
    recordings for many phones. Returns (sched_by_key, tx, rec) where
    sched_by_key is keyed by (phone10, project_norm)."""
    variants = []
    for p in phones10:
        variants += _variants(p)
    sched_by_key = {}
    for s in _client["ASBLVoiceBot"]["call_schedules"].find(
            {"to": {"$in": variants}},
            {"to": 1, "call_outcome": 1, "metadata.project": 1, "called_at": 1, "call_id": 1}):
        key = (_l10(s.get("to")), _norm_proj((s.get("metadata") or {}).get("project")))
        sched_by_key.setdefault(key, []).append(s)
    call_ids = [s.get("call_id") for lst in sched_by_key.values() for s in lst if s.get("call_id")]
    tx, rec = {}, {}
    if call_ids:
        for t in _client["ASBLVoiceBot"]["call_transcripts"].find(
                {"call_sid": {"$in": call_ids}},
                {"call_sid": 1, "duration_secs": 1, "recording_url": 1, "transcript": 1}):
            tx[t.get("call_sid")] = t
        for r in _client["ASBLVoiceBot"]["call_recordings"].find(
                {"call_sid": {"$in": call_ids}},
                {"call_sid": 1, "recording_url": 1, "recording_duration_secs": 1}):
            rec[r.get("call_sid")] = r
    return sched_by_key, tx, rec


def _ai_from(sched, tx, rec):
    """Build the ai_calls block (counts + call_conversations) from a pre-fetched
    schedule list + transcript/recording maps, in the cron's shape."""
    def _turns(raw):
        out = []
        for t in (raw or []):
            spk = str(t.get("speaker") or "")
            is_agent = any(ch.isalpha() for ch in spk)
            out.append({"message": t.get("text", ""),
                        "type": "outbound" if is_agent else "inbound",
                        "speaker": spk if is_agent else "",
                        "date": t.get("ts") or t.get("date")})
        return out

    convs, ans, na, nc, durs = [], 0, 0, 0, []
    for s in sorted(sched, key=lambda c: str(c.get("called_at") or "")):
        sid = s.get("call_id")
        oc = s.get("call_outcome")
        t = tx.get(sid) or {}
        r = rec.get(sid) or {}
        dur = t.get("duration_secs") or r.get("recording_duration_secs")
        if oc in CALL_ANSWERED:        ans += 1
        elif oc in CALL_NOT_ANSWERED:  na += 1
        elif oc in CALL_NOT_CONNECTED: nc += 1
        if dur and float(dur) > 0:
            durs.append(float(dur))
        convs.append({
            "call_sid": sid,
            "call_direction": "outbound",
            "call_outcome": oc,
            "started_at": s.get("called_at"),
            "duration_secs": dur,
            "recording_url": r.get("recording_url") or t.get("recording_url"),
            "recording_duration_secs": r.get("recording_duration_secs"),
            "transcript": _turns(t.get("transcript")),
        })
    total = len(convs)
    return {
        "outbound_calls_triggered": total, "calls_answered": ans,
        "calls_not_answered": na, "calls_not_connected": nc, "inbound_calls": 0,
        "pickup_ratio": round(ans / (ans + na), 2) if (ans + na) else None,
        "average_call_duration_secs": round(sum(durs) / len(durs), 2) if durs else 0.0,
        "call_conversations": convs,
    }


def _last_outcome(ai):
    convs = ai.get("call_conversations") or []
    return convs[-1].get("call_outcome") if convs else None


# ── list ────────────────────────────────────────────────────────────────────
@bulk_bp.route("/api/bulk-calls", methods=["GET"])
def list_bulk():
    rows = []
    for d in _profiles().find({"bulk_manual": True}):
        ai = d.get("ai_calls") or {}
        rows.append({
            "id": d["_id"], "lead_id": d["_id"],
            "phone": d.get("lead_phone_number"), "project": d.get("project"),
            "name": d.get("lead_name"),
            "calls": ai.get("outbound_calls_triggered", 0),
            "answered": ai.get("calls_answered", 0),
            "pickup_ratio": ai.get("pickup_ratio", 0.0),
            "last_outcome": _last_outcome(ai),
            "stage": d.get("stage") or "New Lead",
            "stage_reason": d.get("stage_reason"),
            "flag": d.get("flag"),
            "display_id": d.get("bulk_lead_id"),   # lead id captured from the upload sheet (— if legacy row)
        })
    rows.sort(key=lambda r: (r["project"] or "", r["phone"] or ""))
    return jsonify({"rows": rows, "count": len(rows)})


# ── upload (admin) ───────────────────────────────────────────────────────────
@bulk_bp.route("/api/bulk-calls/upload", methods=["POST"])
def upload_bulk():
    if not _admin():
        return jsonify({"error": "Admin only"}), 403
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "no file uploaded"}), 400
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    except Exception as e:
        return jsonify({"error": f"could not read xlsx: {e}"}), 400
    ws = None
    for cand in wb.worksheets:
        first = next(cand.iter_rows(values_only=True), None) or ()
        if any(str(x).strip().lower() in ("phone", "mobile", "phone number") for x in first if x):
            ws = cand
            break
    ws = ws or wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return jsonify({"error": "empty sheet"}), 400
    hdr = [str(x).strip().lower() if x is not None else "" for x in rows[0]]

    def col(*names):
        for n in names:
            if n in hdr:
                return hdr.index(n)
        return None

    pi = col("phone", "mobile", "phone number")
    pj = col("project")
    ni = col("name", "customer_name", "customer name")
    li = col("lead id", "leadid", "lead_id", "source_lead_id")   # optional: shown in place of phone
    if pi is None:
        return jsonify({"error": "no 'Phone' column found in sheet"}), 400

    # 1. de-dupe (phone, project) within the sheet so each lead is built once
    seen = set()
    items = []
    skipped = 0
    for r in rows[1:]:
        p10 = _l10(r[pi]) if pi < len(r) else ""
        if not p10:
            skipped += 1
            continue
        proj_raw = r[pj] if (pj is not None and pj < len(r)) else ""
        pn = _norm_proj(proj_raw)
        if (p10, pn) in seen:
            continue
        seen.add((p10, pn))
        name = r[ni] if (ni is not None and ni < len(r)) else None
        lead_id_val = r[li] if (li is not None and li < len(r)) else None
        items.append((p10, pn, proj_raw, name, lead_id_val))

    # 2. bulk-fetch all call data (3 queries total — NOT per lead, so a 600-row
    #    sheet doesn't fan out to thousands of round-trips / a serverless timeout).
    sched_by_key, tx, rec = _fetch_call_data([it[0] for it in items])

    # 3. build + upsert each profile from the in-memory data
    now = datetime.now(timezone.utc)
    added = updated = 0
    for p10, pn, proj_raw, name, lead_id_val in items:
        lid = _lead_id(p10, pn)
        ai = _ai_from(sched_by_key.get((p10, pn), []), tx, rec)
        set_fields = {
            "lead_id": lid,  # api_lead(/lead/<id>) resolves by this field
            "lead_phone_number": p10, "project": (str(proj_raw).strip() or None),
            "lead_source": "Bulk Manual Calling", "bulk_manual": True,
            "ai_calls": ai, "updated_at": now, "computed_at": now,
        }
        if name:
            set_fields["lead_name"] = str(name).strip()
        if lead_id_val not in (None, ""):   # only set when the sheet has it — don't wipe on re-upload
            set_fields["bulk_lead_id"] = str(lead_id_val).strip()
        res = _profiles().update_one(
            {"_id": lid},
            {"$set": set_fields,
             "$setOnInsert": {"stage": "New Lead", "stage_reason": None, "flag": None,
                              "website": {}, "whatsapp": {}, "lead_created_at": now}},
            upsert=True)
        if res.upserted_id is not None:
            added += 1
        else:
            updated += 1
    return jsonify({"added": added, "updated": updated, "skipped": skipped,
                    "total": added + updated})


# ── manual stage (admin) ───────────────────────────────────────────────────────
@bulk_bp.route("/api/bulk-calls/stage", methods=["POST"])
def set_bulk_stage():
    if not _admin():
        return jsonify({"error": "Admin only"}), 403
    b = request.get_json(silent=True) or {}
    lead_id = b.get("id", "")
    stage = (b.get("stage") or "").strip()
    if stage not in STAGES:
        return jsonify({"error": f"unknown stage {stage!r}"}), 400
    upd = {"stage": stage, "updated_at": datetime.now(timezone.utc),
           "updated_by": session.get("user_email")}
    upd["stage_reason"] = (b.get("reason") or "Not Responding") if stage == "Not Interested" else None
    res = _profiles().update_one({"_id": lead_id, "bulk_manual": True}, {"$set": upd})
    if not res.matched_count:
        return jsonify({"error": "not found"}), 404
    return jsonify({"ok": True, "stage": stage, "reason": upd["stage_reason"]})


# ── manual flag (admin) ────────────────────────────────────────────────────────
@bulk_bp.route("/api/bulk-calls/flag", methods=["POST"])
def set_bulk_flag():
    if not _admin():
        return jsonify({"error": "Admin only"}), 403
    b = request.get_json(silent=True) or {}
    lead_id = b.get("id", "")
    flag = b.get("flag") or None
    if flag is not None and flag not in FLAGS:
        return jsonify({"error": f"unknown flag {flag!r}"}), 400
    res = _profiles().update_one({"_id": lead_id, "bulk_manual": True},
                                 {"$set": {"flag": flag, "updated_at": datetime.now(timezone.utc)}})
    if not res.matched_count:
        return jsonify({"error": "not found"}), 404
    return jsonify({"ok": True, "flag": flag})
